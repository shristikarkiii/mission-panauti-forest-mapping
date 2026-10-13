"""Step 3 - turn the clipped rasters into forest-change statistics.

Counts, for Panauti Municipality only:

* baseline forest area in 2000 (canopy cover at or above the threshold),
* tree-cover loss for each year 2001-2024,
* tree-cover gain over 2000-2012,
* the resulting cumulative loss and remaining forest.

Areas are accumulated per raster row so that the latitude-dependent size of a
lat/lon pixel is respected. Writes a CSV and an Excel workbook to
``outputs/tables/`` plus a machine-readable summary used by the maps and the
README.

    python src/analyze_forest_change.py
"""

from __future__ import annotations

import csv
import json

import numpy as np
import rasterio

import config
import geo

SUMMARY = config.TABLES / "summary.json"
CSV_PATH = config.TABLES / "panauti_forest_change_2001_2024.csv"
XLSX_PATH = config.TABLES / "panauti_forest_change_2001_2024.xlsx"

COLUMNS = [
    ("year", "Year"),
    ("loss_ha", "Tree cover loss (ha)"),
    ("cumulative_loss_ha", "Cumulative loss (ha)"),
    ("forest_remaining_ha", "Forest remaining (ha)"),
    ("share_of_2000_forest_pct", "Share of 2000 forest (%)"),
    ("annual_loss_rate_pct", "Annual loss rate (%)"),
]


def read_layers() -> tuple[dict[str, np.ndarray], object, tuple[int, int]]:
    layers = {}
    transform = None
    shape = None
    for layer in config.GFC_LAYERS:
        with rasterio.open(config.raster_path(layer)) as src:
            layers[layer] = src.read(1)
            transform, shape = src.transform, src.shape
    return layers, transform, shape


def main() -> None:
    layers, transform, shape = read_layers()
    geometry = geo.load_boundary()

    inside = geo.boundary_mask(geometry, shape, transform)
    land = layers["datamask"] == 1  # 1 = land, 2 = permanent water, 0 = no data
    study = inside & land
    area_ha = np.broadcast_to(geo.pixel_area_ha(transform, shape), shape)

    forest2000 = study & (layers["treecover2000"] >= config.CANOPY_THRESHOLD)
    baseline_ha = float(area_ha[forest2000].sum())
    study_area_ha = float(area_ha[study].sum())

    # Hansen counts loss only where there was tree cover in 2000, so the
    # annual series is restricted to the baseline forest mask.
    lossyear = np.where(forest2000, layers["lossyear"], 0)
    gain_ha = float(area_ha[study & (layers["gain"] == 1)].sum())

    rows = []
    cumulative = 0.0
    for year in range(config.LOSS_YEAR_START, config.LOSS_YEAR_END + 1):
        code = year - 2000
        loss_ha = float(area_ha[lossyear == code].sum())
        standing_before = baseline_ha - cumulative
        cumulative += loss_ha
        rows.append(
            {
                "year": year,
                "loss_ha": round(loss_ha, 2),
                "cumulative_loss_ha": round(cumulative, 2),
                "forest_remaining_ha": round(baseline_ha - cumulative, 2),
                "share_of_2000_forest_pct": round(100 * (baseline_ha - cumulative) / baseline_ha, 3),
                "annual_loss_rate_pct": round(100 * loss_ha / standing_before, 4)
                if standing_before > 0
                else 0.0,
            }
        )

    peak = max(rows, key=lambda r: r["loss_ha"])
    summary = {
        "place": f"{config.PLACE} Municipality, Kavrepalanchok, Nepal",
        "dataset": f"Hansen Global Forest Change {config.GFC_VERSION}",
        "canopy_threshold_pct": config.CANOPY_THRESHOLD,
        "period": [config.LOSS_YEAR_START, config.LOSS_YEAR_END],
        "study_area_ha": round(study_area_ha, 1),
        "forest_2000_ha": round(baseline_ha, 1),
        "forest_2000_pct_of_area": round(100 * baseline_ha / study_area_ha, 1),
        "total_loss_ha": round(cumulative, 1),
        "total_loss_pct_of_2000_forest": round(100 * cumulative / baseline_ha, 2),
        "gain_2000_2012_ha": round(gain_ha, 1),
        "net_change_ha": round(gain_ha - cumulative, 1),
        "mean_annual_loss_ha": round(cumulative / len(rows), 2),
        "peak_loss_year": peak["year"],
        "peak_loss_ha": peak["loss_ha"],
        "forest_2024_ha": round(baseline_ha - cumulative, 1),
    }

    write_csv(rows)
    write_xlsx(rows, summary)
    SUMMARY.write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(f"[ok  ] {CSV_PATH.relative_to(config.ROOT)}")
    print(f"[ok  ] {XLSX_PATH.relative_to(config.ROOT)}")
    print(f"[ok  ] {SUMMARY.relative_to(config.ROOT)}")
    print()
    for key in ("forest_2000_ha", "total_loss_ha", "gain_2000_2012_ha", "net_change_ha",
                "peak_loss_year", "forest_2024_ha"):
        print(f"       {key:<28} {summary[key]}")


def write_csv(rows: list[dict]) -> None:
    with CSV_PATH.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=[k for k, _ in COLUMNS])
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows: list[dict], summary: dict) -> None:
    """Excel copy of the same numbers, for stakeholders who work in Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Annual change"
    ws.append([label for _, label in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in rows:
        ws.append([row[key] for key, _ in COLUMNS])
    ws.freeze_panes = "A2"
    for i, (_, label) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(label) + 2)

    meta = wb.create_sheet("Summary")
    meta.append(["Metric", "Value"])
    meta["A1"].font = meta["B1"].font = Font(bold=True)
    for key, value in summary.items():
        meta.append([key.replace("_", " ").capitalize(), str(value)])
    meta.column_dimensions["A"].width = 34
    meta.column_dimensions["B"].width = 44

    wb.save(XLSX_PATH)


if __name__ == "__main__":
    main()
